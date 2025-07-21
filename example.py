from core.central_dispatcher import CentralDispatcher
from agents.vision_module_agent import VisionModuleAgent
from agents.decision_agent import DecisionAgent
from agents.action_agent import ActionAgent
from agents.speech_agent import SpeechAgent
from agents.memory_agent import MemoryAgent
from agents.profit_seeker_agent import ProfitSeekerAgent
from agents.dashboard_agent import DashboardAgent
from agents.debug_agent import DebugAgent
from agents.emotion_agent import EmotionAgent
from agents.routine_agent import RoutineAgent
from agents.web_agent import WebAgent
from agents.training_agent import TrainingAgent
from agents.screen_analysis_agent import ScreenAnalysisAgent
from core.ai_core import AICore
from logging_config import setup_logging
import logging
import tkinter as tk
from core.dashboard import Dashboard
import threading
import requests
import time
import docx

def main():
    """
    This is an example of how to use the VisionModuleAgent to analyze a Word document.
    """
    setup_logging()
    logging.info("Sentinel AI Online.")

    ai_core = AICore()
    dispatcher = CentralDispatcher()

    decision_agent = DecisionAgent(dispatcher, ai_core)
    action_agent = ActionAgent(dispatcher, ai_core)
    speech_agent = SpeechAgent(dispatcher)
    memory_agent = MemoryAgent(dispatcher)
    vision_module_agent = VisionModuleAgent(dispatcher)

    dispatcher.register_agent("DecisionAgent", decision_agent)
    dispatcher.register_agent("ActionAgent", action_agent)
    dispatcher.register_agent("SpeechAgent", speech_agent)
    dispatcher.register_agent("MemoryAgent", memory_agent)
    dispatcher.register_agent("VisionModuleAgent", vision_module_agent)

    dispatcher.start_agents()

    # Create a sample Word document
    doc = docx.Document()
    doc.add_paragraph("This is an invoice for a contract.")
    doc.add_paragraph("Action Item: Please review the contract and sign it.")
    doc.save("sample.docx")

    # Analyze the document
    vision_module_agent.receive_message({"type": "analyze_document", "path": "sample.docx"})

    # Keep the main thread alive
    while True:
        time.sleep(1)

if __name__ == "__main__":
    main()
